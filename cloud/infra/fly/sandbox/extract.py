"""Print the text of a PDF, docx, or xlsx file. Runs inside nsjail."""

import sys

path = sys.argv[1]
kind = path.rsplit(".", 1)[-1]

if kind == "pdf":
    from pypdf import PdfReader

    for number, page in enumerate(PdfReader(path).pages[:300], 1):
        print(f"--- page {number} ---")
        print(page.extract_text() or "")
elif kind == "docx":
    import docx

    document = docx.Document(path)
    for paragraph in document.paragraphs:
        print(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            print("\t".join(cell.text for cell in row.cells))
elif kind == "xlsx":
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        print(f"--- sheet {sheet.title} ---")
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= 5000:
                print("[more rows not shown]")
                break
            print("\t".join("" if value is None else str(value) for value in row))
else:
    sys.exit(f"unsupported file type: {kind}")
